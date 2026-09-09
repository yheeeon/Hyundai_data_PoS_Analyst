from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

import openpyxl
import pandas as pd

REGISTRATION_PATTERN = "*자동차_등록자료_통계.xlsx"
SALES_PATTERN = "hmc-sales-by-model-*.xlsx"
REGISTRATION_SHEET = "10.연료별_등록현황"
SALES_SHEET = "Unit Sales by Model"
MONTH_NAMES = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}
CATEGORIES = {"PC", "RV", "CV"}
VALID_VEHICLE_TYPES = {"승용", "승합", "화물", "특수"}


def _clean(value):
    if value is None or pd.isna(value):
        return None
    value = str(value).strip()
    return value or None


def _number(value):
    value = _clean(value)
    if value is None or value == "-":
        return None
    number = pd.to_numeric(value.replace(",", ""), errors="coerce")
    return None if pd.isna(number) else int(number)


def _date_from_registration(path: Path) -> tuple[int, int]:
    match = re.search(r"(20\d{2})년[_-]?(\d{1,2})월", path.name)
    if not match:
        raise ValueError(f"기준월을 파일명에서 찾을 수 없습니다: {path.name}")
    return int(match.group(1)), int(match.group(2))


def _date_from_sales(path: Path) -> tuple[int, int]:
    year = re.search(r"y(20\d{2})", path.stem, re.IGNORECASE)
    month = next((number for name, number in MONTH_NAMES.items() if name in path.stem.lower()), None)
    if year is None or month is None:
        raise ValueError(f"현대 판매 파일의 기준월을 해석할 수 없습니다: {path.name}")
    return int(year.group(1)), month


def _rows(path: Path, sheet_name: str) -> list[tuple]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"{path.name}에 '{sheet_name}' 시트가 없습니다.")
    return [tuple(row) for row in workbook[sheet_name].iter_rows(values_only=True)]


def _find_registration_header(rows: list[tuple]) -> int:
    for index, row in enumerate(rows):
        if len(row) > 1 and _clean(row[0]) == "연료별" and _clean(row[1]) == "시도별":
            return index
    raise ValueError("등록 현황 헤더 행을 찾을 수 없습니다.")


def extract_registration(path: Path) -> list[dict]:
    rows = _rows(path, REGISTRATION_SHEET)
    header_index = _find_registration_header(rows)
    header = rows[header_index]
    city_columns = {index: _clean(value) for index, value in enumerate(header) if _clean(value) and _clean(value) != "계"}
    city_columns = {index: city for index, city in city_columns.items() if index >= 3}
    year, month = _date_from_registration(path)
    records = []
    current_fuel = None
    current_vehicle_type = None

    for row in rows[header_index + 2:]:
        fuel = _clean(row[0] if len(row) > 0 else None)
        vehicle_type = _clean(row[1] if len(row) > 1 else None)
        usage = _clean(row[2] if len(row) > 2 else None)
        current_fuel = fuel or current_fuel
        current_vehicle_type = vehicle_type or current_vehicle_type
        if current_vehicle_type not in VALID_VEHICLE_TYPES or usage in {None, "계", "소계"} or current_fuel == "총계":
            continue
        for column, city in city_columns.items():
            if column >= len(row):
                continue
            count = _number(row[column])
            if count is not None:
                records.append({
                    "year": year, "month": month, "city": city, "fuel": current_fuel,
                    "vehicle_type": current_vehicle_type, "usage": usage, "count": count,
                })
    return records


def powertrain_group(fuel: str) -> str:
    if "하이브리드" in fuel:
        return "HEV"
    if fuel == "전기":
        return "EV"
    if fuel in {"수소", "수소전기"}:
        return "FCEV"
    if fuel == "휘발유":
        return "Gasoline"
    if fuel == "경유":
        return "Diesel"
    if fuel in {"엘피지", "LPG"}:
        return "LPG"
    return "Other"


def _find_sales_header(rows: list[tuple]) -> tuple[int, int, int, dict[int, str]]:
    for row_index, row in enumerate(rows):
        month_columns = {}
        for column, value in enumerate(row):
            label = _clean(value)
            if label:
                match = re.match(r"([A-Za-z]{3})", label)
                if match and match.group(1).lower() in MONTH_NAMES:
                    month_columns[column] = match.group(1).title()
        if month_columns:
            model_header_column = next((column for column, value in enumerate(row) if _clean(value) == "Model"), None)
            if model_header_column is not None:
                return row_index, model_header_column, model_header_column + 1, month_columns
    raise ValueError("현대 판매 월별 헤더 행을 찾을 수 없습니다.")


def extract_sales(path: Path) -> list[dict]:
    rows = _rows(path, SALES_SHEET)
    header_index, label_column, model_column, month_columns = _find_sales_header(rows)
    year, source_month = _date_from_sales(path)
    records = []
    section = None
    category = None

    for row in rows[header_index + 1:]:
        label = _clean(row[label_column] if label_column >= 0 and label_column < len(row) else None)
        model = _clean(row[model_column] if model_column < len(row) else None)
        if label in {"Domestic", "Export"}:
            section, category = label, None
            continue
        if label in CATEGORIES:
            category = label
        if section is None or category is None or model is None or model.lower() in {"sub-total", "subtotal", "total", "grand total"}:
            continue
        for column, month_name in month_columns.items():
            units = _number(row[column] if column < len(row) else None)
            if units is not None:
                records.append({
                    "year": year, "source_month": source_month, "month": MONTH_NAMES[month_name.lower()],
                    "market": section, "category": category, "model": model, "units": units,
                })
    return records


def _model_powertrain(model: str) -> str:
    upper = model.upper()
    if any(upper.startswith(prefix) for prefix in ("IONIQ 5", "IONIQ 6", "IONIQ 9", "GV60")):
        return "EV"
    if "PHEV" in upper:
        return "PHEV"
    if "FCEV" in upper or "NEXO" in upper:
        return "FCEV"
    if "HEV" in upper:
        return "HEV"
    if "EV" in upper:
        return "EV"
    return "ICE"


def _model_family(model: str) -> str:
    family = re.sub(r"\s*\(.*\)", "", model)
    family = re.sub(r"\s+(EV|HEV|PHEV|N)$", "", family).strip()
    return "Casper" if family == "Casper EV" else family


def _dimension(values: Iterable[str], key: str, label: str) -> pd.DataFrame:
    unique = sorted({_clean(value) for value in values if _clean(value) is not None})
    return pd.DataFrame({key: range(1, len(unique) + 1), label: unique})


def build_dataset(data_dir: str | Path = "data", output_dir: str | Path = "output/warehouse") -> dict[str, pd.DataFrame]:
    data_path = Path(data_dir)
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    registration_files = sorted(data_path.glob(REGISTRATION_PATTERN))
    sales_files = sorted(data_path.glob(SALES_PATTERN))
    if not registration_files:
        raise FileNotFoundError(f"등록 XLSX를 찾을 수 없습니다: {data_path}")
    if not sales_files:
        raise FileNotFoundError(f"현대 판매 XLSX를 찾을 수 없습니다: {data_path}")

    registration = pd.DataFrame(record for path in registration_files for record in extract_registration(path))
    sales = pd.DataFrame(record for path in sales_files for record in extract_sales(path))
    registration["powertrain"] = registration["fuel"].map(powertrain_group)
    sales["powertrain"] = sales["model"].map(_model_powertrain)
    sales["model_family"] = sales["model"].map(_model_family)
    registration["date_id"] = registration["year"] * 100 + registration["month"]
    sales["date_id"] = sales["year"] * 100 + sales["source_month"]

    dates = pd.DataFrame(sorted(set(registration.date_id) | set(sales.date_id)), columns=["date_id"])
    dates["year"] = dates.date_id // 100
    dates["month"] = dates.date_id % 100
    dates["year_month"] = dates.date_id.astype(str)
    regions = _dimension(registration.city, "region_id", "region_name")
    vehicle_types = _dimension(registration.vehicle_type, "vehicle_type_id", "vehicle_type_name")
    fuels = _dimension(registration.fuel, "fuel_id", "fuel_name").merge(
        registration[["fuel", "powertrain"]].drop_duplicates().rename(columns={"fuel": "fuel_name"}), on="fuel_name", how="left"
    )
    powertrains = _dimension(set(registration.powertrain) | set(sales.powertrain), "powertrain_id", "powertrain_name")
    models = _dimension(sales.model, "model_id", "model_name").merge(
        sales[["model", "model_family", "category", "powertrain"]].drop_duplicates().rename(columns={"model": "model_name"}), on="model_name", how="left"
    )

    registration = registration.merge(regions, left_on="city", right_on="region_name").merge(vehicle_types, left_on="vehicle_type", right_on="vehicle_type_name").merge(fuels[["fuel_id", "fuel_name"]], left_on="fuel", right_on="fuel_name")
    registration = registration[["date_id", "region_id", "fuel_id", "vehicle_type_id", "usage", "count"]].rename(columns={"count": "registered_qty"})
    sales = sales.merge(models[["model_id", "model_name"]], left_on="model", right_on="model_name")
    sales = sales[["date_id", "model_id", "market", "category", "month", "units"]].rename(columns={"month": "sales_month", "units": "sales_qty"})

    tables = {
        "dim_date": dates, "dim_region": regions, "dim_vehicle_type": vehicle_types,
        "dim_fuel": fuels, "dim_powertrain": powertrains, "dim_model": models,
        "fact_registration": registration, "fact_sales": sales,
    }
    for name, frame in tables.items():
        frame.to_csv(output_path / f"{name}.csv", index=False, encoding="utf-8-sig")
    return tables
